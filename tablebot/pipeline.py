"""逐题处理、并发调度、结果落盘、生成提交文件。"""
import json
import re
import shutil
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

from .config import Config
from .llm import QwenClient
from .normalize import finalize
from .prompts import build_messages
from .render import load_document


# 账户类错误（欠费、Key 无效、无权限）：后面的题必然全部失败，立刻停止，不再发请求
_ACCOUNT_ERRORS = ("Arrearage", "InvalidApiKey", "invalid_api_key", "AccessDenied", "401",
                   "FreeTierOnly", "AllocationQuota")
_abort = threading.Event()


def model_tag(model: str, thinking: bool) -> str:
    return re.sub(r"[^A-Za-z0-9._-]", "_", model) + ("_think" if thinking else "")


def _migrate_flat_results(base: Path) -> None:
    """旧版本把所有模型的结果平铺在 work/results/ 下，这里按模型挪进各自的子目录。"""
    for p in list(base.glob("*.json")):
        try:
            r = json.loads(p.read_text(encoding="utf-8"))
            d = base / model_tag(r.get("model", "unknown"), bool(r.get("thinking")))
            d.mkdir(parents=True, exist_ok=True)
            shutil.move(str(p), str(d / p.name))
        except Exception:  # noqa: BLE001
            pass


def results_dir(cfg: Config, tag: Optional[str] = None) -> Path:
    base = cfg.work_dir / "results"
    base.mkdir(parents=True, exist_ok=True)
    _migrate_flat_results(base)
    d = base / (tag or cfg.result_tag or model_tag(cfg.model, cfg.enable_thinking))
    d.mkdir(parents=True, exist_ok=True)
    return d


def solve_one(q: dict, cfg: Config, client: Optional[QwenClient], dry_run: bool = False) -> dict:
    rec = {"id": q["id"], "file_name": q["file_name"], "question_type": q["question_type"],
           "answer_format": q["answer_format"], "question": q["question"],
           "table_hint": q["table_hint"], "fixes": q["fixes"], "model": cfg.model,
           "thinking": cfg.enable_thinking, "answer": "", "error": None}
    if _abort.is_set():
        rec["error"] = "已中止：账户不可用（见上方报错）"
        rec["aborted"] = True
        return rec
    try:
        if not q["path"]:
            raise FileNotFoundError(f"找不到文件 {q['file_name']}")
        doc = load_document(q["path"], cfg)
        messages = build_messages(q, doc)
        rec["pages"] = len(doc.pages)
        rec["has_text_layer"] = bool(doc.text)
        if dry_run:
            chars = sum(len(c.get("text", "")) for c in messages[1]["content"] if c["type"] == "text")
            # 千问 VL 约每 28x28 像素 1 个 token；中文约 1 字 1 token
            img_tok = sum(min(16384, -(-w // 28) * -(-h // 28)) for w, h in doc.page_sizes)
            rec["est_tokens"] = img_tok + chars + len(messages[0]["content"])
            return rec
        resp = client.chat(messages)
        rec["model_used"] = resp.get("model", cfg.model)
        rec["raw"] = resp["content"]
        rec["usage"] = resp.get("usage", {})
        rec["cached"] = resp.get("cached", False)
        try:
            out = finalize(q, resp["content"], cfg.array_numeric_as_number, cfg.trust_expression)
        except ValueError:
            # 输出不是合法 JSON：把原输出发回去，让模型改成合法 JSON（重试一次）
            retry = messages + [
                {"role": "assistant", "content": resp["content"]},
                {"role": "user", "content": "你上面的输出不是合法 JSON（可能有未转义的引号或被截断）。"
                                            "请重新只输出一个合法的 JSON 对象，内容精简，不要有任何其他文字。"}]
            resp = client.chat(retry)
            rec["raw_retry"] = resp["content"]
            out = finalize(q, resp["content"], cfg.array_numeric_as_number, cfg.trust_expression)
            out["notes"].append("首次输出不是合法 JSON，已重试")
        # 非结构题答案为空：提醒模型再答一次
        if q["question_type"] != "structure" and out["answer"] in ("", "[]"):
            retry = messages + [
                {"role": "assistant", "content": resp["content"]},
                {"role": "user", "content": "answer 不能为空。如果题目问的是文件/图片本身的属性（是否……），"
                                            "请直接观察图片回答“是”或“否”；否则请找最相近的字段给出最可能的答案。"
                                            "只输出 JSON。"}]
            try:
                r2 = client.chat(retry)
                out2 = finalize(q, r2["content"], cfg.array_numeric_as_number, cfg.trust_expression)
                if out2["answer"] not in ("", "[]"):
                    rec["raw_retry"] = r2["content"]
                    out2["notes"].append("首次答案为空，已重试")
                    out = out2
            except Exception:  # noqa: BLE001
                pass
        rec["answer"] = out["answer"]
        rec["notes"] = out["notes"]

        # 多次采样投票（结构题不投票）
        if cfg.votes > 1 and q["question_type"] != "structure":
            answers = [out["answer"]]
            for k in range(1, cfg.votes):
                try:
                    r2 = client.chat(messages, temperature=cfg.vote_temperature, sample=k)
                    answers.append(finalize(q, r2["content"], cfg.array_numeric_as_number,
                                            cfg.trust_expression)["answer"])
                except Exception:  # noqa: BLE001
                    continue
            best = _majority(answers)
            rec["votes"] = answers
            if best != out["answer"]:
                rec["notes"].append(f"投票改为 {best}（首答 {out['answer']}）")
                rec["answer"] = best
    except Exception as e:  # noqa: BLE001
        rec["error"] = f"{type(e).__name__}: {e}"
        if any(k in rec["error"] for k in _ACCOUNT_ERRORS):
            _abort.set()
    return rec


def _majority(answers: List[str]) -> str:
    """多数票；平票时取第一个（温度 0 的那次）。空答案不参与。"""
    from collections import Counter
    valid = [a for a in answers if a]
    if not valid:
        return answers[0] if answers else ""
    cnt = Counter(valid)
    top = max(cnt.values())
    for a in valid:  # 按出现顺序，保证平票取首答
        if cnt[a] == top:
            return a
    return valid[0]


def run(questions: List[dict], cfg: Config, use_cache: bool = True, dry_run: bool = False,
        workers: Optional[int] = None) -> List[dict]:
    client = None if dry_run else QwenClient(cfg, use_cache=use_cache)
    rdir = results_dir(cfg)
    workers = workers or cfg.workers
    done, t0 = [], time.time()
    tok_in = tok_out = 0
    # 先在主线程里把所有文件渲染一遍（结果缓存在内存），并发阶段就不会再碰 pdfium
    paths = sorted({q["path"] for q in questions if q["path"]})
    for k, pth in enumerate(paths, 1):
        print(f"\r预处理文件 {k}/{len(paths)}", end="", flush=True)
        try:
            load_document(pth, cfg)
        except Exception as e:  # noqa: BLE001
            print(f"\n  文件 {pth} 预处理失败：{e}")
    print()
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futs = {ex.submit(solve_one, q, cfg, client, dry_run): q for q in questions}
        for i, f in enumerate(as_completed(futs), 1):
            rec = f.result()
            done.append(rec)
            if not dry_run and not rec.get("aborted"):
                fp = rdir / f"{rec['id']}.json"
                keep_old = False
                if rec["error"] and fp.exists():
                    try:
                        keep_old = not json.loads(fp.read_text(encoding="utf-8")).get("error")
                    except Exception:  # noqa: BLE001
                        pass
                if not keep_old:
                    fp.write_text(json.dumps(rec, ensure_ascii=False, indent=1), encoding="utf-8")
                u = rec.get("usage") or {}
                if not rec.get("cached"):
                    tok_in += u.get("prompt_tokens", 0) or 0
                    tok_out += u.get("completion_tokens", 0) or 0
            status = "错误 " + rec["error"][:80] if rec["error"] else (rec["answer"][:60] or "(空)")
            flag = "C" if rec.get("cached") else " "
            print(f"[{i}/{len(questions)}] {flag} id={rec['id']:<4} {rec['question_type']:<9} {status}",
                  flush=True)
    dt = time.time() - t0
    n_err = sum(1 for r in done if r["error"])
    print(f"\n完成 {len(done)} 题，用时 {dt:.0f}s，失败 {n_err} 题；本次新消耗 tokens：输入 {tok_in}，输出 {tok_out}")
    if _abort.is_set():
        print("!! 账户不可用（欠费 / Key 无效 / 无权限），已停止发送请求。处理好账户后，"
              "用同样的命令加 --only-missing 继续，已完成的题不会重复付费。")
    return sorted(done, key=lambda r: r["id"])


def refinalize(questions: List[dict], cfg: Config) -> None:
    """不调用 API：用当前的清洗规则，从保存的模型原始输出重新生成所有答案（改了 normalize.py 后用）。"""
    qmap = {q["id"]: q for q in questions}
    rdir = results_dir(cfg)
    changed = total = 0
    for fp in rdir.glob("*.json"):
        r = json.loads(fp.read_text(encoding="utf-8"))
        q = qmap.get(int(r["id"]))
        if not q or r.get("error") or not r.get("raw") or r.get("votes"):
            continue  # 失败的、没有原始输出的、投票得出的，保持不变
        total += 1
        out = None
        for content in (r.get("raw_retry"), r["raw"]):
            if not content:
                continue
            try:
                o = finalize(q, content, cfg.array_numeric_as_number, cfg.trust_expression)
                if o["answer"] not in ("", "[]"):
                    out = o
                    break
                out = out or o
            except Exception:  # noqa: BLE001
                continue
        if out is None:
            continue
        if out["answer"] != r.get("answer"):
            changed += 1
            print(f"  id={r['id']}: {str(r.get('answer'))[:50]}  →  {out['answer'][:50]}")
        r["answer"], r["notes"] = out["answer"], out["notes"]
        fp.write_text(json.dumps(r, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"重新清洗 {total} 题，答案有变化的 {changed} 题（{rdir.name}）")


def load_results(cfg: Config, tag: Optional[str] = None) -> Dict[int, dict]:
    out = {}
    for p in results_dir(cfg, tag).glob("*.json"):
        try:
            r = json.loads(p.read_text(encoding="utf-8"))
            out[int(r["id"])] = r
        except Exception:  # noqa: BLE001
            pass
    return out


def _fallback_answer(q: dict) -> str:
    """评测系统不接受空单元格（Excel 无法保存空字符串），空答案用占位值兜底。"""
    fmt = q["answer_format"]
    if q["question_type"] == "structure" or fmt == "json":
        return '{"row_count": 0, "col_count": 0, "cells": []}'
    if fmt == "json_array":
        return "[]"
    if fmt == "number":
        return "0"
    return "否" if "是否" in q["question"] else "无"


def validate_submission(path: Path, questions: List[dict]) -> List[str]:
    """上传前自检：题号齐全且不重复、无空单元格、json / json_array 答案能解析。"""
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    rows = list(ws.iter_rows(values_only=True))
    probs = []
    if rows[0][:2] != ("id", "answer"):
        probs.append(f"表头应为 id, answer，实际 {rows[0]}")
    body = rows[1:]
    want = [q["id"] for q in questions]
    got = [r[0] for r in body]
    if got != want:
        probs.append(f"题号不一致：应 {len(want)} 行，实际 {len(got)} 行")
    fmt = {q["id"]: q for q in questions}
    for i, a in body:
        if a is None or str(a).strip() == "":
            probs.append(f"id={i} 答案为空")
            continue
        q = fmt.get(i)
        if q and (q["answer_format"] in ("json", "json_array") or q["question_type"] == "structure"):
            try:
                json.loads(a)
            except Exception:  # noqa: BLE001
                probs.append(f"id={i} 不是合法 JSON：{str(a)[:40]}")
    return probs


def write_submission(questions: List[dict], cfg: Config, name: Optional[str] = None,
                     fallbacks: Optional[List[str]] = None) -> Path:
    """按题目清单顺序写 submission.xlsx：每题一行。当前模型没有结果的题，依次用 fallbacks
    （其他模型的结果目录名）补；仍然没有的用占位值兜底。answer 一律存为文本。"""
    from openpyxl import Workbook

    res = load_results(cfg)
    used_fb = {}
    for tag in fallbacks or []:
        fb = load_results(cfg, tag)
        for qid, r in fb.items():
            cur = res.get(qid)
            if (not cur or cur.get("error") or not cur.get("answer")) and r.get("answer") and not r.get("error"):
                res[qid] = r
                used_fb[tag] = used_fb.get(tag, 0) + 1
    cfg.out_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["id", "answer"])
    missing = []
    for q in questions:
        r = res.get(q["id"])
        ans = r["answer"] if r and not r.get("error") else ""
        if not str(ans).strip():
            missing.append(q["id"])
            ans = _fallback_answer(q)
        ws.append([q["id"], ans])
        ws.cell(row=ws.max_row, column=2).number_format = "@"  # 文本格式，防止 Excel 改写数字
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    path = cfg.out_dir / (name or f"submission_{stamp}.xlsx")
    wb.save(path)
    wb.save(cfg.out_dir / "submission.xlsx")
    print(f"已写出 {path}（同时覆盖 outputs/submission.xlsx）；共 {len(questions)} 题，"
          f"其中 {cfg.result_tag or model_tag(cfg.model, cfg.enable_thinking)} 的答案 {len(questions) - sum(used_fb.values()) - len(missing)} 题"
          + "".join(f"，{t} 补 {n} 题" for t, n in used_fb.items()))
    problems = validate_submission(path, questions)
    if problems:
        print("!! 提交文件自检发现问题，先别上传：")
        for p_ in problems[:20]:
            print("   ", p_)
    else:
        print("提交文件自检通过：行数、题号、无空单元格、JSON 格式均正常，可以上传。")
    if missing:
        print(f"注意：{len(missing)} 题没有答案，已填占位值：{missing[:30]}")
    return path
