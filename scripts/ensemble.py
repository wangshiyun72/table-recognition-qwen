"""
多模型投票：用第三个模型当“裁判”，在两个主模型答案不一致的题上决定取哪一个。

规则（逐题）：
  - 主模型 A 与 B 答案相同 → 直接用；
  - 不同时，裁判 C 与谁一致就取谁；
  - 三者都不同 → 取 A（通常是最强的模型）。

用法（在项目根目录执行）：
  python scripts/ensemble.py --a outputs/A.xlsx --b outputs/B.xlsx --referee outputs/C.xlsx \
                             --out outputs/submission_ensemble.xlsx

本项目 80.5 分的版本即由此生成：
  A = qwen3.8-max（思考模式）全量答案
  B = 80.1 分版本（qwen3.7-plus + qwen3.8-max）
  C = qwen3-vl-plus 版本（71.9 分）
"""
import argparse
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from tablebot.config import Config  # noqa: E402
from tablebot.data import load_questions  # noqa: E402
from tablebot.pipeline import validate_submission  # noqa: E402


def load(path: str) -> dict:
    rows = list(load_workbook(path).active.iter_rows(values_only=True))[1:]
    return {int(r[0]): r[1] for r in rows if r and r[0] is not None}


def main():
    p = argparse.ArgumentParser(description="三模型投票合并提交文件")
    p.add_argument("--a", required=True, help="主模型 A 的提交文件（三者都不同时取它）")
    p.add_argument("--b", required=True, help="主模型 B 的提交文件")
    p.add_argument("--referee", required=True, help="裁判模型 C 的提交文件")
    p.add_argument("--out", default="outputs/submission_ensemble.xlsx")
    args = p.parse_args()

    cfg = Config()
    questions = load_questions(cfg.tests_path, cfg.files_dir)
    a, b, c = load(args.a), load(args.b), load(args.referee)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["id", "answer"])
    stat = {"一致": 0, "裁判选A": 0, "裁判选B": 0, "三者不同取A": 0}
    for q in questions:
        i = q["id"]
        if a.get(i) == b.get(i):
            ans, k = a.get(i), "一致"
        elif c.get(i) == b.get(i):
            ans, k = b.get(i), "裁判选B"
        elif c.get(i) == a.get(i):
            ans, k = a.get(i), "裁判选A"
        else:
            ans, k = a.get(i), "三者不同取A"
        stat[k] += 1
        ws.append([i, ans])
        ws.cell(row=ws.max_row, column=2).number_format = "@"

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"已写出 {out}：{stat}")
    probs = validate_submission(out, questions)
    print("自检：" + ("通过" if not probs else "; ".join(probs[:10])))


if __name__ == "__main__":
    main()
